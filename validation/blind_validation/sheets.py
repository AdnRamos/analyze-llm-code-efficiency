"""
Fase 3 (Revisor A, A1): prepara as planilhas de classificação manual CEGA.

Lê validation/reports/triage/amostra_fase3.csv (casos escolhidos pela triagem) e validation/reports/triage/triagem_submissoes.csv,
reconstrói o código submetido com as mesmas funções de validation/triage.py e gera:

  validation/blind_validation/fase3_classificacao_adenilson.xlsx   } idênticas; sem sugestão automática, sem flags
  validation/blind_validation/fase3_classificacao_rodrigo.xlsx     } e sem o nome do modelo (classificação cega)
  validation/blind_validation/fase3_mapa_cego.csv                  id_cego -> chave, sugestão automática, flags (NÃO abrir antes de classificar)

Lotes:
  nucleo_kappa  : 200 casos classificados pelos DOIS autores (base do kappa de Cohen).
                  Todos os casos com sugestão C, 70 com sugestão B e o restante com sugestão A, sorteados.
                  A sobreamostragem de B e C evita um kappa degenerado pela prevalência de A.
  restante      : demais casos da amostra.

Uso: python validation/blind_validation/preparar_fase3.py     (requer openpyxl; semente fixa)
"""
import csv
import os
import random
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
VALID = os.path.dirname(HERE)
sys.path.insert(0, VALID)
import triage as triagem  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

SEED = 20261013
N_NUCLEO = 200
N_B_NUCLEO = 70
MAX_CEL = 32000  # limite do Excel é 32.767 caracteres por célula
csv.field_size_limit(10 ** 9)


def ler_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def corta(s):
    s = s or ""
    return s if len(s) <= MAX_CEL else s[:MAX_CEL] + "\n[... cortado: %d caracteres no total]" % len(s)


def assinatura(spec):
    ms = ", ".join("%s(%d)" % (m, k) for m, k in spec["methods"].items())
    return "class %s: %s" % (spec["cls"], ms)


def main():
    amostra = ler_csv(os.path.join(VALID, "reports", "triage", "amostra_fase3.csv"))
    tri = {"%s|%s|%s" % (r["modelo"], r["linguagem"], r["id"]): r
           for r in ler_csv(os.path.join(VALID, "reports", "triage", "triagem_submissoes.csv"))}
    qs, specs, gen, subs = triagem.carregar()
    sub = {(s["modelo"], s["linguagem"], int(s["id_questao"])): s for s in subs}

    rng = random.Random(SEED)
    por_sug = {"A": [], "B": [], "C": []}
    for r in amostra:
        por_sug.setdefault(r["sugestao"], []).append(r)
    for lst in por_sug.values():
        lst.sort(key=lambda r: r["chave"])
    nucleo = list(por_sug["C"])
    nucleo += rng.sample(por_sug["B"], min(N_B_NUCLEO, len(por_sug["B"])))
    resto_a = N_NUCLEO - len(nucleo)
    nucleo += rng.sample(por_sug["A"], min(resto_a, len(por_sug["A"])))
    chaves_nucleo = {r["chave"] for r in nucleo}
    restante = [r for r in amostra if r["chave"] not in chaves_nucleo]
    rng.shuffle(nucleo)
    rng.shuffle(restante)

    linhas, mapa = [], []
    for i, (lote, r) in enumerate([("nucleo_kappa", x) for x in nucleo] + [("restante", x) for x in restante], 1):
        modelo, ling, qid = r["chave"].split("|")
        key = (modelo.replace("/", "_"), ling, int(qid))
        g, s = gen[key], sub[key]
        recon = triagem.reconstruir(g, s.get("ts_epoch"))
        recon = recon[0] if isinstance(recon, tuple) else recon
        t = tri[r["chave"]]
        erro = s.get("erro_runtime") or s.get("erro_compilacao") or ""
        id_cego = "F3-%04d" % i
        linhas.append([id_cego, lote, s["slug"], assinatura(specs[int(qid)]["Python3"]), s["status"],
                       corta(erro), corta(g.get("resposta")), corta(recon or ""), "", "", ""])
        mapa.append([id_cego, lote, r["chave"], r["sugestao"], r["regra_sugestao"], r["grupo"],
                     r["flags_principais"], t["tipo_erro"]])

    for autor in ("adenilson", "rodrigo"):
        escrever_xlsx(os.path.join(HERE, "fase3_classificacao_%s.xlsx" % autor), autor, linhas)
    with open(os.path.join(HERE, "fase3_mapa_cego.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["id_cego", "lote", "chave", "sugestao_auto", "regra_sugestao", "grupo_amostra",
                    "flags_principais", "tipo_erro"])
        w.writerows(mapa)
    comp = {k: sum(1 for r in nucleo if r["sugestao"] == k) for k in ("A", "B", "C")}
    print("núcleo kappa: %d (A=%d, B=%d, C=%d); restante: %d; total: %d"
          % (len(nucleo), comp["A"], comp["B"], comp["C"], len(restante), len(linhas)))


INSTRUCOES = [
    ("Classificação manual — Revisor A, ponto A1 (Fase 3)", None),
    ("", None),
    ("O que fazer", "Para cada linha da aba 'Classificar', leia a resposta bruta do modelo, o código que foi "
     "submetido e o erro do LeetCode, e preencha APENAS as colunas amarelas: 'categoria' (lista A/B/C/?) e, "
     "se quiser, 'observacao'. Comece pelo lote 'nucleo_kappa' (as 200 primeiras linhas)."),
    ("Classificação cega", "A planilha não mostra o modelo, a sugestão automática nem as flags da triagem, "
     "para não influenciar a decisão. Não consulte o colega nem o arquivo fase3_mapa_cego.csv antes de "
     "terminar o núcleo."),
    ("A — Artefato do pipeline", "A resposta bruta contém um programa aderente (classe e métodos esperados, "
     "com a aridade certa), e o código submetido deixou de ser aderente ou falha por algo introduzido pela "
     "extração ou pelo wrapper (classe perdida, resto de cerca ``` ou de rótulo, classe de design envolvida "
     "em Solution, recuo quebrado)."),
    ("B — Descumprimento de formato", "Nenhuma parte da resposta bruta é um programa aderente: classe, método "
     "ou aridade errados, só raciocínio, rótulos ou texto dentro do próprio bloco de código, ou testes, print "
     "ou input() deixados pelo modelo que quebram a execução."),
    ("C — Erro genuíno", "O código submetido é o que o modelo escreveu, é aderente e falha em execução "
     "(IndexError, TypeError interno, recursão etc.)."),
    ("? — Não sei", "Use quando o caso for misto ou não der para decidir (ex.: o modelo descumpriu o formato "
     "E o extrator piorou). Explique na observação."),
    ("Assinatura esperada", "A coluna 'assinatura_esperada' vem do starter da questão: classe e métodos com o "
     "número de parâmetros posicionais (sem self). Nomes de parâmetros e anotações de tipo podem diferir; "
     "métodos auxiliares a mais são permitidos."),
    ("Exemplo (já classificado)", "Questão vowel-strings-in-range (2586). A resposta bruta traz o bloco "
     "'class Solution: def vowelStrings(self, words, left, right)' duas vezes. O código submetido é "
     "'class Solution:\\n    solution:' e o erro é SyntaxError na linha 2 → categoria A (o programa do "
     "modelo era válido; o extrator pegou o texto errado)."),
    ("Depois", "Salve o arquivo com o mesmo nome. O script validation/blind_validation/kappa.py lê as duas planilhas, calcula o "
     "kappa de Cohen no núcleo e lista as discordâncias para consenso."),
]

CAB = ["id_cego", "lote", "questao", "assinatura_esperada", "veredito", "erro_leetcode",
       "resposta_bruta_modelo", "codigo_submetido", "categoria", "observacao", "tempo_min_opcional"]
LARG = [10, 13, 26, 34, 14, 50, 80, 70, 11, 40, 10]


def escrever_xlsx(path, autor, linhas):
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Instruções"
    fonte = Font(name="Arial", size=10)
    for i, (a, b) in enumerate(INSTRUCOES, 1):
        ws0.cell(i, 1, a).font = Font(name="Arial", size=12 if i == 1 else 10, bold=True)
        if b:
            c = ws0.cell(i, 2, b)
            c.font = fonte
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws0.cell(len(INSTRUCOES) + 2, 1, "Classificador").font = Font(name="Arial", bold=True)
    ws0.cell(len(INSTRUCOES) + 2, 2, autor.capitalize()).font = fonte
    ws0.column_dimensions["A"].width = 30
    ws0.column_dimensions["B"].width = 110

    ws = wb.create_sheet("Classificar")
    amarelo = PatternFill("solid", start_color="FFFF00")
    for j, h in enumerate(CAB, 1):
        c = ws.cell(1, j, h)
        c.font = Font(name="Arial", bold=True)
        if h in ("categoria", "observacao", "tempo_min_opcional"):
            c.fill = amarelo
        ws.column_dimensions[c.column_letter].width = LARG[j - 1]
    for i, row in enumerate(linhas, 2):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font = Font(name="Consolas" if CAB[j - 1] in ("erro_leetcode", "resposta_bruta_modelo",
                                                             "codigo_submetido") else "Arial", size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if CAB[j - 1] in ("categoria", "observacao", "tempo_min_opcional"):
                c.fill = amarelo
        ws.row_dimensions[i].height = 150
    dv = DataValidation(type="list", formula1='"A,B,C,?"', allow_blank=True,
                        error="Use A, B, C ou ?", errorTitle="Categoria inválida", showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add("I2:I%d" % (len(linhas) + 1))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:%s%d" % (ws.cell(1, len(CAB)).column_letter, len(linhas) + 1)
    wb.save(path)


if __name__ == "__main__":
    main()
